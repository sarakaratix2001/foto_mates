from __future__ import annotations

import argparse
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
import csv
import json
from datetime import datetime, timezone
import math
import os
from pathlib import Path
import queue
import re
import sqlite3
import sys
import threading
import time
from evaluation import GROUPS as EVALUATION_GROUPS, SCORE_THRESHOLD, evaluate, lines as evaluation_lines
from extract_readings import (BACKENDS as MODEL_BACKENDS, backend as model_backend, OCR_MODEL, PHOTO_TYPES, RETAKE_SCORE, WORKERS_PER_SLOT, photo_retake_reason, _reading_digits, ocr_image,
                              server_slots, extract_blank, extract_no_meter, low_score_reasons, score_reading)

NAME = re.compile(r'^(?P<contract>[^_]+)_(?P<stamp>\d{14})_(?P<operator>[^_]*)_(?P<reading>\d+)__'
                  r'(?P<register>[^_]+)')
AUDIT_DB = Path(__file__).resolve().parent / os.environ.get('FOTOMATES_AUDIT_DB', 'audit_readings.sqlite3')
BATCH_LOOKUP = 900  # below SQLite's limit of host parameters per statement


def processing_exclusions():
    path = Path(__file__).resolve().parent / 'processing_exclusions.json'
    if not path.exists():
        return set()
    return {row['name'] for row in json.loads(path.read_text(encoding='utf-8'))['files']}

SCHEMA = """
CREATE TABLE IF NOT EXISTS readings (
    picture_name      TEXT PRIMARY KEY,
    contract_id       TEXT,
    date              TEXT,     -- YYYY-MM-DD the photo was taken
    month             TEXT,     -- YYYY-MM
    register          TEXT,     -- ExpAktiv, ExpReaktiv, ExpPik, ...
    saved_reading     TEXT,     -- the reading in the system (from the file name)
    extracted_reading TEXT,     -- every digit read from the photo, leading zeros and decimals included
    extracted_whole   TEXT,     -- the consumption read: no leading zeros, decimals stripped
    extracted_decimals INTEGER, -- how many digits were stripped as decimals
    match             TEXT NOT NULL,  -- yes / no
    status            TEXT NOT NULL,  -- done / missing / no_meter / blank / no_saved_reading / error
    score             REAL,
    retake            INTEGER,
    seconds           REAL,
    error             TEXT,
    read_at           TEXT NOT NULL,
    model             TEXT,     -- the model that read the photo
    visual_decimals   INTEGER,  -- visually identified mechanical fractional wheels, or NULL
    clear             REAL,     -- the model's chance of QUALITY: CLEAR, 0..1 (information only)
    quality           TEXT      -- the model's QUALITY answer, CLEAR or POOR (information only)
) WITHOUT ROWID;  -- the rows live in the picture_name index itself: one copy, smaller file
CREATE INDEX IF NOT EXISTS readings_month ON readings(month, status, match);
CREATE INDEX IF NOT EXISTS readings_contract ON readings(contract_id, month, register);
"""

# Everything else known about a read, kept so a result can be checked or a rule applied
# again later without reading the photo again. Added to old databases on connect.
DETAIL_COLUMNS = (
    ('operator', 'TEXT'),               # who took the photo, from the file name
    ('latitude', 'REAL'),               # where, from the file name (41-35034332 -> 41.35034332)
    ('longitude', 'REAL'),
    ('device_type', 'TEXT'),            # digital / mechanical, as the model answered TYPE
    ('digit_probabilities', 'TEXT'),    # one per digit of extracted_reading, e.g. "0.982 1.000 0.931"
    ('second_choices', 'TEXT'),         # per digit, the likeliest other digit: "7:0.120 -:- 0:0.010"
    ('weakest_digit', 'REAL'),          # the lowest of digit_probabilities
    ('code_removed', 'TEXT'),           # register code cut from the front of READING (380, 1580, ...)
    ('code_probabilities', 'TEXT'),     # one per digit of code_removed
    ('reject_reason', 'TEXT'),          # why retake is 1
    ('too_many_digits', 'INTEGER'),     # 1: 12 or more digits read, score set to 0
    ('rolling', 'TEXT'),                # mechanical wheels caught turning, e.g. "6:6-7"
    ('reread', 'INTEGER'),              # 1: read a second time, enlarged, after an unsure first read
    ('photo_sent', 'TEXT'),             # what the model saw: original / enlarged(2x) / enhanced(is_dark,...)
    ('image_width', 'INTEGER'),         # the photo as stored, in pixels
    ('image_height', 'INTEGER'),
    ('blur', 'REAL'),                   # image checks before the read (higher blur value: sharper)
    ('contrast', 'REAL'),
    ('brightness', 'REAL'),
    ('glare_fraction', 'REAL'),
    ('image_flags', 'TEXT'),            # is_blurry, is_dark, has_glare, ... that were true
    ('model_answer', 'TEXT'),           # the model's whole answer, with the timing lines after it
)


def connect(path, **options):
    db = sqlite3.connect(path, **options)
    db.execute('PRAGMA journal_mode=WAL')
    db.execute('PRAGMA synchronous=FULL')
    db.executescript(SCHEMA)
    present = {column[1] for column in db.execute('PRAGMA table_info(readings)')}
    added = (('model', 'TEXT'), ('extracted_decimals', 'INTEGER'), ('visual_decimals', 'INTEGER'),
             ('clear', 'REAL'), ('quality', 'TEXT'), *DETAIL_COLUMNS)
    for name, kind in added:
        if name not in present:  # a database from before the column
            db.execute(f'ALTER TABLE readings ADD COLUMN {name} {kind}')
    attempts = {column[1] for column in db.execute('PRAGMA table_info(reading_attempts)')}
    for name, kind in added:
        if attempts and name not in attempts:  # the copy of earlier results gets the column too
            db.execute(f'ALTER TABLE reading_attempts ADD COLUMN {name} {kind}')
    # Preserve the previous result of every image when a retry replaces the
    # current row. The trigger also covers writers already running older code.
    db.execute('CREATE TABLE IF NOT EXISTS reading_attempts AS SELECT * FROM readings')
    columns = ','.join(COLUMNS)
    values = ','.join('NEW.' + column for column in COLUMNS)
    # Made again on every connect, so a database from before a new column also records it.
    db.execute('DROP TRIGGER IF EXISTS readings_record_attempt')
    db.execute(f'''CREATE TRIGGER readings_record_attempt
        AFTER INSERT ON readings BEGIN
        INSERT INTO reading_attempts ({columns}) VALUES ({values}); END''')
    db.execute('CREATE INDEX IF NOT EXISTS reading_attempts_picture ON reading_attempts(picture_name, read_at)')
    db.commit()
    return db


MAX_DECIMALS = 3  # the most decimals a display shows (000000.123)


def split_decimals(reading, saved, visual_decimals=None):
    saved = re.sub(r'\D', '', saved or '')
    reading = re.sub(r'\D', '', reading or '')  # a reading stored with its decimal point
    if not reading:
        return None, None, 'no'
    if visual_decimals is not None:
        if type(visual_decimals) is not int or not 0 <= visual_decimals < len(reading):
            raise ValueError('visual_decimals must be a nonnegative count smaller than the digit count')
        whole = str(int(reading) // 10 ** visual_decimals)
        return whole, visual_decimals, 'yes' if saved and int(whole) == int(saved) else 'no'
    if not saved:
        return reading.lstrip('0') or '0', None, 'no'
    saved_value = int(saved)
    # A saved 0 has no whole digits: its 0 is one of the display's leading zeros, so what
    # follows them is decimals (000000.19 reads 0.19, not 1.9), at most MAX_DECIMALS of them.
    saved_digits = math.floor(math.log10(saved_value)) + 1 if saved_value else 0
    leading_zeros = len(reading) - len(reading.lstrip('0'))
    decimals = max(0, len(reading) - leading_zeros - saved_digits)
    if not saved_value:
        decimals = min(decimals, MAX_DECIMALS)
    whole = int(reading) // 10 ** decimals
    return str(whole), decimals, 'yes' if whole == saved_value else 'no'


def parse_name(name, month=None):

    match = NAME.match(name)
    if not match:
        return dict(contract_id=name.split('_')[0] if '_' in name else Path(name).stem,
                    date=None, month=month, register=None, saved_reading=None)
    try:
        taken = datetime.strptime(match['stamp'], '%Y%m%d%H%M%S')
        date, month = taken.strftime('%Y-%m-%d'), taken.strftime('%Y-%m')
    except ValueError:
        date = None
    return dict(contract_id=match['contract'], date=date, month=month,
                register=match['register'], saved_reading=match['reading'])


def photos(folders):
    stack = [Path(folder) for folder in folders]
    excluded = processing_exclusions()
    while stack:
        folder = stack.pop()
        found = 0
        print(f'Source scan: listing {folder} (a day folder on the share takes a few minutes) ...', flush=True)
        # Each photo is yielded as soon as it is listed, so reading starts at once.
        # Read through prefetch(), the listing never waits for the reads: the SMB
        # enumeration stays open only as long as the listing itself takes.
        with os.scandir(folder) as entries:
            for entry in entries:
                if entry.is_dir(follow_symlinks=False):
                    stack.append(Path(entry.path))
                elif entry.name not in excluded and Path(entry.name).suffix.lower() in PHOTO_TYPES:
                    found += 1
                    if found % 10000 == 0:
                        print(f'Source scan: {folder}: {found} photos listed so far ...', flush=True)
                    yield Path(entry.path)
        print(f'Source scan: {folder}: {found} eligible photos', flush=True)


def prefetch(iterable, size=0):
    """The items of iterable, produced in a thread of their own: folders on the
    share are listed while their photos are being read, so the model server is
    not left idle for the minutes a 100 000-file listing takes. The queue has no
    limit (size 0): the listing never waits for the reads, so a network folder's
    enumeration is never held open. An error of the listing is raised here,
    where the caller handles it."""
    items, end = queue.Queue(size), object()

    def fill():
        try:
            for item in iterable:
                items.put(item)
        except BaseException as exc:
            items.put(exc)
        finally:
            items.put(end)

    threading.Thread(target=fill, name='folder-listing', daemon=True).start()
    while (item := items.get()) is not end:
        if isinstance(item, BaseException):
            raise item
        yield item


def unread(db, paths, retry_errors=True, only_failed=False):
    """The paths whose picture is not yet in the database (or failed before);
    with only_failed, just the ones that failed before."""
    batch = []
    excluded = processing_exclusions()

    def flush():
        names = [path.name for path in batch]
        marks = ','.join('?' * len(names))
        if only_failed:
            failed = {row[0] for row in db.execute(
                f"SELECT picture_name FROM readings WHERE picture_name IN ({marks}) AND status = 'error'", names)}
            fresh = [path for path in batch if path.name in failed]
            batch.clear()
            return fresh
        skip = "AND status != 'error'" if retry_errors else ''
        known = {row[0] for row in db.execute(
            f'SELECT picture_name FROM readings WHERE picture_name IN ({marks}) {skip}', names)}
        fresh = [path for path in batch if path.name not in known]
        batch.clear()
        return fresh

    for path in paths:
        if path.name in excluded:
            continue
        batch.append(path)
        if len(batch) >= BATCH_LOOKUP:
            yield from flush()
    if batch:
        yield from flush()


def failed_photos(db, folders):
    """The photos saved with a technical error (status 'error', "Gabim teknik" in
    the reports), found by name inside the given folders instead of by listing
    them: a day folder on the share holds over 100 000 files, and only these few
    are needed. A photo whose date (from its name) matches a folder's YYYY/MM/DD
    goes to that folder unchecked (a file check on the share costs ~2 s; a photo
    missing there fails its read and stays a technical error); any other is
    looked for in every folder, and one not found is reported and left as it is."""
    excluded = processing_exclusions()
    roots = [Path(folder) for folder in folders]
    missing = 0
    for (name,) in db.execute("SELECT picture_name FROM readings WHERE status = 'error' ORDER BY picture_name"):
        if name in excluded:
            continue
        # The date opens the name's second field, also in names NAME does not fit
        # (a longer time stamp, empty fields).
        match = re.match(r'[^_]+_(\d{4})(\d{2})(\d{2})', name)
        day = Path(*match.groups()) if match else None
        dated = next((root / name for root in roots if day and root.parts[-3:] == day.parts), None)
        path = dated or next((root / name for root in roots if (root / name).is_file()), None)
        if path:
            yield path
        else:
            missing += 1
            print(f'not found in the given folders, left as a technical error: {name}', flush=True)
    if missing:
        print(f'{missing} failed photo(s) not found in the given folders', flush=True)


def csv_photos(filename):
    """Explicit retry paths, without scanning folders or reading a saved database."""
    seen = set()
    excluded = processing_exclusions()
    with Path(filename).open(encoding='utf-8-sig', newline='') as source:
        reader = csv.DictReader(source)
        if not {'picture_name', 'source_path'}.issubset(reader.fieldnames or []):
            raise ValueError('Retry CSV requires picture_name and source_path columns')
        for line, row in enumerate(reader, 2):
            path = Path(row['source_path'])
            name = row['picture_name']
            if not row['source_path'] or not path.is_absolute() or path.name != name:
                raise ValueError(f'Invalid source_path or picture_name on retry CSV line {line}')
            if path.suffix.lower() not in PHOTO_TYPES:
                raise ValueError(f'Unsupported photo extension on retry CSV line {line}')
            if name in seen or name in excluded:
                continue
            seen.add(name)
            yield path


def record_retry_failure(filename, row):
    """Persist failed paths independently of the optional in-memory database."""
    filename = Path(filename)
    filename.parent.mkdir(parents=True, exist_ok=True)
    fields = ['picture_name', 'source_path', 'reason', 'error', 'read_at', 'status']
    header = not filename.exists() or filename.stat().st_size == 0
    with filename.open('a', encoding='utf-8-sig' if header else 'utf-8', newline='') as output:
        writer = csv.DictWriter(output, fields)
        if header:
            writer.writeheader()
        writer.writerow({key: ('technical_error' if key == 'reason' else row.get(key, '')) for key in fields})
        output.flush()
        os.fsync(output.fileno())


REREAD_BELOW = float(os.environ.get('FOTOMATES_REREAD_BELOW', '0.7'))
# Every small photo is already read enlarged to FOTOMATES_MIN_SIDE (2x for 240x400), so the reread goes further.
REREAD_SCALE = float(os.environ.get('FOTOMATES_REREAD_SCALE', '3'))


def read_scored(path, saved, model=None, scale=1.0):
    """One read of a photo with its consumption and its one score."""
    result = ocr_image(path, model=model, scale=scale)
    cells = _reading_digits(result.text) or ''
    # Scored from the photo alone, then compared with the saved reading.
    parts = score_reading(cells, result.digit_confidence, result.decimal_digits, result.digit_margin,
                          result.digit_runner_up, result.digit_alternatives,
                          result.device_type == 'mechanical') if cells else None
    whole, decimals, match = split_decimals(cells or None, saved, result.decimal_digits)
    return dict(result=result, cells=cells, whole=whole, decimals=decimals, match=match, parts=parts,
                score=parts['total'] if parts else None)


def read_best(path, saved, model=None):
    """The read of a photo: as it is, and once more enlarged only when the first
    is unsure (score below REREAD_BELOW, or no reading); the read with the
    higher score is kept. Visual mechanical decimal evidence overrides the
    legacy saved-length rule; the saved value never decides which read is kept. Adds 'reread' (whether
    a second read was made) and 'seconds' (model time of every read made)."""
    first = read_scored(path, saved, model)
    first.update(reread=False, seconds=first['result'].reading_seconds or 0.0)
    if REREAD_SCALE == 1:
        return first  # no resizing: the same image read again gives the same answer (temperature 0)
    if first['score'] is not None and first['score'] >= REREAD_BELOW:
        return first
    if extract_no_meter(first['result'].text):
        return first  # no meter in the photo: enlarging it shows none either
    try:
        second = read_scored(path, saved, model, REREAD_SCALE)
    except Exception:
        return first  # the first read stands when the second fails
    seconds = (first['result'].reading_seconds or 0.0) + (second['result'].reading_seconds or 0.0)
    best = second if (second['score'] or -1) > (first['score'] or -1) else first
    return dict(best, reread=True, seconds=seconds)


_PLACE = re.compile(r'_(\d+)-(\d+)_(\d+)-(\d+)\.[^.]+$')


def name_details(name):
    """The operator and the coordinates the file name carries (…_OPERATOR_…_41-35034332_19-86586937.jpg)."""
    match, place = NAME.match(name), _PLACE.search(name)
    lat = lon = None
    if place and place[1] != '0':
        lat, lon = float(f'{place[1]}.{place[2]}'), float(f'{place[3]}.{place[4]}')
    return dict(operator=(match['operator'] or None) if match else None, latitude=lat, longitude=lon)


def read_details(result, reading, reread):
    """What a read knows beyond the reading: per-digit probabilities and second choices,
    the register code cut off, the model's answer and what it was shown."""
    def numbers(values):
        return ' '.join(f'{p:.3f}' for p in values) if values else None
    probabilities = list(result.digit_confidence or [])
    aligned = reading is not None and len(probabilities) == len(reading)
    second = []
    for index, options in enumerate(result.digit_alternatives or [] if aligned else []):
        others = sorted(((p, d) for d, p in (options or {}).items() if d != reading[index]), reverse=True)
        second.append(f'{others[0][1]}:{others[0][0]:.3f}' if others else '-:-')
    image = result.image_report or {}
    flags = [name for name, value in image.items() if name.startswith(('is_', 'has_')) and value]
    answer = re.findall(r'READING:[ \t]*([^\n]*)', result.text, re.IGNORECASE)
    written = re.sub(r'\D', '', answer[-1]) if answer else ''
    return dict(
        device_type=result.device_type,
        digit_probabilities=numbers(probabilities) if aligned else None,
        second_choices=' '.join(second) or None,
        weakest_digit=round(min(probabilities), 3) if aligned and probabilities else None,
        code_removed=written[:len(result.code_confidence or [])] or None,
        code_probabilities=numbers(result.code_confidence),
        rolling=', '.join(f"{w['position']}:{w['lower']}-{w['upper']}" if w.get('position') else 'turning'
                          for w in result.rolling or []) or None,
        reread=int(bool(reread)), photo_sent=result.evidence_source,
        image_width=image.get('width'), image_height=image.get('height'),
        blur=image.get('blur_variance'), contrast=image.get('contrast'), brightness=image.get('brightness'),
        glare_fraction=image.get('glare_fraction'), image_flags=','.join(flags) or None,
        model_answer=result.text)


def read_one(path, month, model=None):
    row = dict(parse_name(path.name, month), **name_details(path.name))
    started = time.monotonic()
    try:
        read = read_best(path, row['saved_reading'], model)
        result = read['result']
    except Exception as exc:
        return dict(row, model=model_backend(model).model, picture_name=path.name, extracted_reading=None, extracted_whole=None,
                    extracted_decimals=None, match='no', status='error', score=None, retake=None,
                    seconds=round(time.monotonic() - started, 2), error=str(exc)[:2000], source_path=str(path.absolute()))
    reading = _reading_digits(result.text) or None  # as read: leading zeros kept
    if extract_blank(result.text):
        status = 'blank'
    elif extract_no_meter(result.text):
        status = 'no_meter'
    elif not reading:
        status = 'missing'
    elif not row['saved_reading']:
        status = 'no_saved_reading'
    else:
        status = 'done'
    whole, decimals, match, score = read['whole'], read['decimals'], read['match'], read['score']
    quality = re.findall(r'^QUALITY:[ \t]*\**[ \t]*(CLEAR|POOR)', result.text, re.IGNORECASE | re.MULTILINE)
    reason = photo_retake_reason(reading, result.clear) if reading else None
    if not reason and score is not None and score < RETAKE_SCORE:
        reason = low_score_reasons(read['parts'])[0] if read['parts'] else 'the overall confidence is low'
    details = read_details(result, reading, read.get('reread'))
    too_long = bool(read['parts'] and read['parts'].get('too_long'))
    return dict(row, **details, model=model_backend(model).model, picture_name=path.name, extracted_reading=reading,
                extracted_whole=whole, extracted_decimals=decimals, match=match,
                status=status, score=score, visual_decimals=result.decimal_digits,
                clear=None if result.clear is None else round(result.clear, 3),
                quality=quality[-1].upper() if quality else None,
                retake=int(bool(reason)), reject_reason=reason, too_many_digits=int(too_long),
                # the same, under the names the run's log lines use
                code=details['code_removed'], reason=reason, too_long=too_long,
                seconds=round(read['seconds'] or time.monotonic() - started, 2), error=None)


COLUMNS = ('picture_name', 'contract_id', 'date', 'month', 'register', 'saved_reading',
           'extracted_reading', 'extracted_whole', 'extracted_decimals', 'match', 'status', 'score',
           'retake', 'seconds', 'error', 'read_at', 'model', 'visual_decimals', 'clear', 'quality',
           *(name for name, _ in DETAIL_COLUMNS))


def save(db, row):
    row['read_at'] = datetime.now(timezone.utc).isoformat(timespec='seconds')
    row['model'] = row.get('model') or OCR_MODEL
    db.execute(f'INSERT OR REPLACE INTO readings ({",".join(COLUMNS)}) '
               f'VALUES ({",".join("?" * len(COLUMNS))})', [row.get(c) for c in COLUMNS])


def command_read(args):
    db = connect(args.db)
    choice = model_backend(args.model)
    slots = server_slots(choice.servers)
    workers = args.workers or (choice.workers if choice.fixed_workers or not slots else slots * WORKERS_PER_SLOT)
    print(f'reading with {workers} worker(s) into {args.db}; model {choice.label}; model server slots: '
          f'{slots or "not reported"}', flush=True)
    started = time.monotonic()
    count = matched = failed = rejected = codes = too_long = 0
    saved_names_so_far = {row[0] for row in db.execute('SELECT picture_name FROM readings')}
    print(f'Resuming with {len(saved_names_so_far)} saved photos; completed photos are skipped.', flush=True)
    model_seconds = 0.0
    pending = set()
    seen = []  
    interruption = None
    last_report_count = 0

    progress_dir = Path(args.word_dir) / 'progress'
    failure_csv = progress_dir / 'retry_failures.csv'
    print(f'Technical failures are saved immediately to: {failure_csv}', flush=True)
    # Checkpoint reports are written in a thread of their own, with their own
    # database connection (WAL lets it read while results are being written):
    # exporting the whole database (seconds) no longer stops photos being sent.
    reporter = ThreadPoolExecutor(max_workers=1, thread_name_prefix='reports')
    report_job = None

    def write_reports(connection):
        try:
            export_reports(connection, progress_dir, args.month, full=True, quiet=True)
            print(f'Reports saved: {progress_dir} (all saved results in this database for the selected month)', flush=True)
        except Exception as exc:
            print(f'Report export could not finish: {exc}. Results remain saved in {args.db}.', flush=True)

    def reader_connection():
        """A connection the export thread can read from. A file database is
        opened again; an in-memory one (--db :memory:) exists only inside this
        connection, so it is copied instead."""
        if str(args.db) in (':memory:', '') or str(args.db).startswith('file::memory'):
            copy = sqlite3.connect(':memory:', check_same_thread=False)
            db.backup(copy)
            return copy
        return sqlite3.connect(args.db, check_same_thread=False)

    def background_reports(reader):
        try:
            write_reports(reader)
        finally:
            reader.close()

    def checkpoint_reports():
        nonlocal last_report_count, report_job
        if report_job is not None and not report_job.done():
            return  # the previous export is still writing; a later checkpoint catches up
        last_report_count = count
        try:
            reader = reader_connection()
        except Exception as exc:
            print(f'Report export could not start: {exc}. Results remain saved in {args.db}.', flush=True)
            return
        report_job = reporter.submit(background_reports, reader)

    def listed(paths):
        for path in paths:
            if not args.no_word:
                seen.append(path.name)
            yield path

    def collect(done):
        nonlocal count, matched, failed, rejected, codes, too_long, model_seconds
        for future in done:
            row = future.result()
            save(db, row)
            db.commit()  # a printed result is already durable, even after a disconnect
            if row['status'] == 'error':
                record_retry_failure(failure_csv, row)
            saved_names_so_far.add(row['picture_name'])
            count += 1
            matched += row['match'] == 'yes'
            failed += row['status'] == 'error'
            rejected += bool(row.get('retake'))
            codes += bool(row.get('code'))
            too_long += bool(row.get('too_long'))
            model_seconds += row['seconds'] or 0
            if args.verbose:
                result = (('MATCH' if row['match'] == 'yes' else 'MISMATCH') if row['status'] == 'done' else
                          {'no_meter': 'NO_METER_FOUND'}.get(row['status'], row['status'].upper()))
                score = '-' if row['score'] is None else f"{row['score']:.3f}"
                elapsed = time.monotonic() - started
                print(f"[{count} this run | {len(saved_names_so_far)} saved] {row['picture_name']} "
                      f"| saved by system={row['saved_reading'] or '-'} "
                      f"| extracted by AI={row['extracted_whole'] or '-'} "
                      f"| raw digits={row['extracted_reading'] or '-'} | score={score} "
                      f"| {result} | photo={row['seconds']}s | elapsed={elapsed:.0f}s "
                      f"| visual decimals={row['visual_decimals'] if row.get('visual_decimals') is not None else '-'}"
                      + (f" | quality={row['quality'] or '-'} ({100 * row['clear']:.0f}% clear, info only)"
                         if row.get('clear') is not None else '')
                      + (f" | code removed={row['code']}" if row.get('code') else '')
                      + (f" | REJECTED: {row['reason']}" if row.get('reason') else ''), flush=True)
                if row['error']:
                    print(f"  error: {row['error']}", flush=True)
            if count % args.progress == 0:
                elapsed = time.monotonic() - started
                rate = count / elapsed
                # failed: photos the model server could not answer at all (not mismatches),
                # shown only when there are any
                print(f'PROGRESS {len(saved_names_so_far)} total saved; {count} read this run in {elapsed / 60:.0f} min, '
                      f'{rate:.2f} photos/s ({rate * 86400:,.0f}/day), {model_seconds / count:.1f} s per photo | '
                      f'{100 * matched / count:.1f}% match, {rejected} rejected ({100 * rejected / count:.1f}%), '
                      f'{codes} register codes removed, {too_long} with 12+ digits'
                      + (f' | {failed} failed (model server gave no answer; --only-failed reads them again)'
                         if failed else ''), flush=True)
        report_every = getattr(args, 'report_every', 1000)
        if report_every and count - last_report_count >= report_every:
            checkpoint_reports()

    with ThreadPoolExecutor(max_workers=workers) as pool:
        try:
            # --only-failed: the photos with a technical error, found by name; no
            # listing of the share, and nothing new is read.
            if getattr(args, 'files_csv', None):
                print(f'Reading only explicit paths in: {args.files_csv}', flush=True)
                paths = unread(db, listed(csv_photos(args.files_csv)), retry_errors=not args.skip_errors)
            else:
                paths = (listed(failed_photos(db, args.folders)) if args.only_failed else
                         unread(db, listed(prefetch(photos(args.folders))), retry_errors=not args.skip_errors))
            for path in paths:
                # Save and print what has finished as it finishes: with many workers
                # the queue is rarely full, so waiting for that would hold every result.
                done, pending = wait(pending, timeout=0 if len(pending) < workers * 2 else None,
                                     return_when=FIRST_COMPLETED)
                collect(done)
                pending.add(pool.submit(read_one, path, args.month, choice.key))
                if args.limit and count + len(pending) >= args.limit:
                    break
        except (OSError, KeyboardInterrupt) as exc:
            interruption = exc
            print(f'Scan interrupted: {exc or "keyboard interrupt"}. Saving remaining completed requests before exit.', flush=True)
        finally:
            while pending:
                done, pending = wait(pending, return_when=FIRST_COMPLETED)
                collect(done)
    db.commit()
    reporter.shutdown(wait=True)  # a checkpoint still writing finishes before the final one
    write_reports(db)
    elapsed = time.monotonic() - started
    print(f'{"stopped safely" if interruption else "done"}: {count} photos in {elapsed:.0f} s'
          + (f' ({count / elapsed:.2f} photos/s)' if elapsed else ''), flush=True)
    if not args.no_word:
        try:
            where, params = ('WHERE month IS ?', [_month(args.month)]) if args.month else ('', [])
            saved_names = [row[0] for row in db.execute(f'SELECT picture_name FROM readings {where}', params)]
            print(f'Word report: {write_word_report(db, args.folders, saved_names, count, args.word_dir)}', flush=True)
        except Exception as exc:
            print(f'Word export could not finish: {exc}. CSV reports and database results are retained.', flush=True)
    db.close()
    if interruption:
        raise SystemExit(1)


def folder_rows(db, names):
    """The stored result of each named photo, as dicts; a photo not read yet is left out."""
    rows = []
    for start in range(0, len(names), BATCH_LOOKUP):
        batch = names[start:start + BATCH_LOOKUP]
        cursor = db.execute(
            'SELECT picture_name AS name, status, saved_reading AS saved, extracted_whole AS extracted, '
            'extracted_reading AS all_digits, score, error FROM readings '
            f'WHERE picture_name IN ({",".join("?" * len(batch))})', batch)
        columns = [column[0] for column in cursor.description]
        rows += [dict(zip(columns, row)) for row in cursor]
    return rows


# Files listed per group in the Word report; the CSV beside it lists them all.
WORD_LIST_MAX = int(os.environ.get('FOTOMATES_WORD_LIST_MAX', '2000'))


def write_word_report(db, folders, names, read_now, out_dir):
    from docx import Document  # python-docx; only this report needs it

    rows = folder_rows(db, sorted(set(names)))
    groups = evaluate(dict(row, no_meter=row['status'] == 'no_meter', blank=row['status'] == 'blank') for row in rows)
    errors = {row['name']: row['error'] for row in rows if row['status'] == 'error'}
    now = datetime.now()

    document = Document()
    document.add_heading('FotoMates reading report', level=1)
    document.add_paragraph(f'Date: {now:%Y-%m-%d %H:%M}\nFolder(s): {", ".join(map(str, folders))}\n'
                           f'Model: {OCR_MODEL}\nThe consumption read from each photo (no leading zeros, '
                           'decimals stripped) is compared with the reading saved in the system.')

    table = document.add_table(rows=1, cols=3, style='Light Grid Accent 1')
    for cell, text in zip(table.rows[0].cells, ('', 'Photos', 'Share / note')):
        cell.text = text
    summary_lines = [('Photos in the folder(s)', len(set(names)), ''),
                     ('Photos read', len(rows), f'{read_now} now, {len(rows) - read_now} in an earlier run')]
    summary_lines += [(label, int(count), f'{100 * int(count) / len(rows):.1f}%' if rows else '-')
                      for _, label, count in evaluation_lines(groups)]
    for label, number, note in summary_lines:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = label, str(number), note

   
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    lists = out / f'fotomates_report_{now:%Y%m%d_%H%M%S}_files.csv'
    with open(lists, 'w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.writer(handle)
        writer.writerow(['group', 'file_name', 'saved_reading', 'extracted_rules_applied', 'all_digits_read',
                         'score', 'error'])
        for key, label in EVALUATION_GROUPS:
            for row in groups['files'][key]:
                writer.writerow([label, row['name'], row['saved'], row['extracted'], row['all_digits'],
                                 '' if row['score'] is None else f'{row["score"]:.3f}', errors.get(row['name'], '')])
    document.add_paragraph(f'Every photo with its group, in full: {lists.name}')

    for key, label in EVALUATION_GROUPS:
        listed = groups['files'][key]
        if key == 'matched' or not listed:
            continue
        document.add_heading(f'{label} ({len(listed)})', level=2)
        if len(listed) > WORD_LIST_MAX:
            document.add_paragraph(f'The first {WORD_LIST_MAX} of {len(listed)}; all of them are in {lists.name}.')
            listed = listed[:WORD_LIST_MAX]
        if key == 'failed':
            for row in listed:
                document.add_paragraph(f'{row["name"]}: {errors.get(row["name"]) or ""}', style='List Bullet')
            continue
        table = document.add_table(rows=1, cols=5, style='Light Grid Accent 1')
        for cell, text in zip(table.rows[0].cells, ('File name', 'Saved reading', 'Extracted (rules applied)',
                                                    'All digits read', 'Score')):
            cell.text = text
        for row in listed:
            cells = table.add_row().cells
            for cell, value in zip(cells, (row['name'], row['saved'], row['extracted'], row['all_digits'],
                                           None if row['score'] is None else f'{row["score"]:.2f}')):
                cell.text = value or ''

    path = out / f'fotomates_report_{now:%Y%m%d_%H%M%S}.docx'
    document.save(path)
    return path


def summary(db, month=None):
    """One row per month: photos, how many were read correctly, mismatched,
    missing or failed, the share correct, contracts, and seconds per photo."""
    where, params = ('WHERE month IS ?', [_month(month)]) if month else ('', [])
    return db.execute(f"""
        SELECT COALESCE(month, 'unknown') AS month,
               COUNT(*) AS photos,
               SUM(match = 'yes') AS read_correctly,
               SUM(status = 'done' AND match = 'no') AS mismatch,
               SUM(status = 'done' AND match = 'no' AND score >= {SCORE_THRESHOLD}) AS mismatch_high_score,
               SUM(status = 'done' AND match = 'no' AND score < {SCORE_THRESHOLD}) AS mismatch_low_score,
               SUM(status = 'done' AND match = 'no' AND score IS NULL) AS mismatch_no_score,
               SUM(status = 'missing') AS missing_reading,
               SUM(status = 'no_meter') AS no_meter,
               SUM(status = 'blank') AS blank,
               SUM(status = 'no_saved_reading') AS no_saved_reading,
               SUM(status = 'error') AS failed,
               ROUND(100.0 * SUM(match = 'yes') / NULLIF(SUM(status IN ('done', 'missing')), 0), 2)
                   AS percent_correct,
               COUNT(DISTINCT contract_id) AS contracts,
               ROUND(AVG(seconds), 2) AS seconds_per_photo
        FROM readings {where} GROUP BY month ORDER BY month""", params)


LIST_KINDS = ('mismatches', 'missing', 'duplicates', 'all')


def _month(month):
    """'unknown' is how photos without a date in their name are listed."""
    return None if month == 'unknown' else month
# The reading in the system next to the consumption read from the photo, then
# every digit as read and how many of them were stripped as decimals.
_PICKED = ('picture_name, contract_id, date, register, saved_reading, extracted_whole AS read_reading, '
           'match, extracted_reading AS all_digits_read, extracted_decimals AS decimals_stripped, '
           'status, score, seconds, model, clear AS clear_info_only, quality AS quality_info_only')


def list_query(kind, month=None, contract=None):
    """The SQL and parameters of one list: mismatches (both readings present,
    different), missing (no reading, no saved reading, or failed), duplicates
    (a contract and register photographed more than once in a month) or all."""
    if kind not in LIST_KINDS:
        raise ValueError(f'kind must be one of {", ".join(LIST_KINDS)}')
    filters, params = [], []
    if month:
        filters.append('month IS ?')
        params.append(_month(month))
    if contract:
        filters.append('contract_id = ?')
        params.append(contract)
    also = ''.join(' AND ' + f for f in filters)
    if kind == 'duplicates':
        inner = ''.join(' AND ' + f for f in filters)
        return (f"""
            SELECT r.contract_id, r.month, r.register, d.photos, r.picture_name, r.date,
                   r.saved_reading, r.extracted_whole AS read_reading, r.match,
                   r.extracted_reading AS all_digits_read
            FROM readings r JOIN (
                SELECT contract_id, month, register, COUNT(*) AS photos FROM readings
                WHERE contract_id IS NOT NULL {inner}
                GROUP BY contract_id, month, register HAVING COUNT(*) > 1) d
              ON r.contract_id = d.contract_id AND r.month IS d.month AND r.register IS d.register
            ORDER BY r.contract_id, r.month, r.register, r.date""", params)
    if kind == 'mismatches':
        return (f"SELECT {_PICKED} FROM readings WHERE status = 'done' AND match = 'no'{also} "
                'ORDER BY month, contract_id', params)
    if kind == 'missing':
        return (f"SELECT {_PICKED}, error FROM readings "
                f"WHERE status IN ('missing', 'no_meter', 'blank', 'error', 'no_saved_reading'){also} "
                'ORDER BY status, month, contract_id', params)
    where = ' WHERE ' + ' AND '.join(filters) if filters else ''
    return (f'SELECT month, {_PICKED}, retake, error, read_at FROM readings{where} '
            'ORDER BY month, contract_id', params)


# One result line of `read --verbose`.
_LOG_LINE = re.compile(r'^(?:\d{4}-\d\d-\d\d \d\d:\d\d:\d\d )?\[\d+(?: this run \| \d+ saved)?\] (?P<name>\S+) \| saved by system=(?P<saved>\S+) \| extracted by AI=\S+ '
                       r'\| raw digits=(?P<raw>\S+) \| score=(?P<score>\S+) \| (?P<result>[A-Z_]+) '
                       r'\| photo=(?P<seconds>[\d.]+)s')


def command_import_log(args):
    db = connect(args.db)
    known = {name for (name,) in db.execute('SELECT picture_name FROM readings')}
    added = failed = present = 0
    for log in args.logs:
        for line in Path(log).read_text(encoding='utf-8', errors='replace').splitlines():
            m = _LOG_LINE.match(line.strip())
            if not m:
                continue
            name, result = m['name'], m['result']
            if result == 'ERROR':
                failed += 1
                continue
            if name in known:
                present += 1
                continue
            row = parse_name(name, args.month)
            raw = None if m['raw'] == '-' else m['raw']
            visual = re.search(r'\| visual decimals=(\d+|-)', line)
            visual_decimals = int(visual[1]) if visual and visual[1] != '-' else None
            whole, decimals, match = split_decimals(raw, row['saved_reading'], visual_decimals)
            status = {'MATCH': 'done', 'MISMATCH': 'done', 'NO_METER_FOUND': 'no_meter'}.get(result, result.lower())
            if status == 'done' and not row['saved_reading']:
                status = 'no_saved_reading'
            score = None if m['score'] == '-' else float(m['score'])
            save(db, dict(row, picture_name=name, extracted_reading=raw, extracted_whole=whole,
                          extracted_decimals=decimals, visual_decimals=visual_decimals, match=match, status=status, score=score,
                          retake=int(score is not None and score < RETAKE_SCORE),
                          seconds=float(m['seconds']), error=None))
            known.add(name)
            added += 1
    db.commit()
    print(f'{added} photo(s) imported into {args.db}; {present} already there; '
          f'{failed} failed in the log and left to read again', flush=True)


def command_rematch(args):
    """Match every stored reading again by the current rule, without reading
    any photo again."""
    db = connect(args.db)
    changed = total = 0
    for name, reading, saved, before, visual_decimals in db.execute(
            'SELECT picture_name, extracted_reading, saved_reading, match, visual_decimals FROM readings').fetchall():
        whole, decimals, match = split_decimals(reading, saved, visual_decimals)
        db.execute('UPDATE readings SET extracted_whole = ?, extracted_decimals = ?, match = ? '
                   'WHERE picture_name = ?', (whole, decimals, match, name))
        total += 1
        changed += match != before
    db.commit()
    print(f'{total} readings matched again, {changed} changed', flush=True)


def write_csv(path, cursor):
    path = Path(path)
    temporary = path.with_name(path.name + '.tmp')
    try:
        with temporary.open('w', newline='', encoding='utf-8-sig') as handle:
            writer = csv.writer(handle)
            writer.writerow([column[0] for column in cursor.description])
            rows = 0
            for row in cursor:
                writer.writerow(row)
                rows += 1
            handle.flush()
            os.fsync(handle.fileno())
        temporary.replace(path)  # retain the last good report if an export fails
        return rows
    finally:
        temporary.unlink(missing_ok=True)


def export_reports(db, out, month=None, full=False, quiet=False):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    write_csv(out / 'summary.csv', summary(db, month))
    comparison_filter = 'WHERE month IS ?' if month else ''
    write_csv(out / 'comparison.csv', db.execute(f'''
        SELECT picture_name AS "File name", saved_reading AS "Saved by system",
               extracted_whole AS "Extracted by AI", extracted_reading AS "Raw digits",
               extracted_decimals AS "Digits stripped", score AS "Score",
               CASE WHEN status = 'error' THEN 'Processing error'
                    WHEN status = 'no_meter' THEN 'No meter'
                    WHEN status = 'blank' THEN 'Blank photo'
                    WHEN status = 'missing' THEN 'No readable digits'
                    WHEN status = 'no_saved_reading' THEN 'No saved reading'
                    WHEN match = 'yes' THEN 'Match'
                    WHEN score >= {SCORE_THRESHOLD} THEN 'Mismatch - high score'
                    WHEN score < {SCORE_THRESHOLD} THEN 'Mismatch - low score'
                    ELSE 'Mismatch - no score' END AS "Result",
               model AS "AI model", error AS "Error",
               CASE WHEN retake = 1 THEN 'yes' WHEN retake = 0 THEN 'no' END AS "Rejected",
               reject_reason AS "Why rejected", device_type AS "Meter type",
               digit_probabilities AS "Digit probabilities", second_choices AS "Second choices",
               weakest_digit AS "Weakest digit", code_removed AS "Code removed",
               code_probabilities AS "Code probabilities", too_many_digits AS "12+ digits",
               reread AS "Read twice", photo_sent AS "Photo sent to model",
               clear AS "Chance of CLEAR (info only)", quality AS "Quality (info only)",
               operator AS "Operator", latitude AS "Latitude", longitude AS "Longitude"
        FROM readings {comparison_filter} ORDER BY month, picture_name
    ''', [_month(month)] if month else []))
    counts = {kind: write_csv(out / f'{kind}.csv', db.execute(*list_query(kind, month)))
              for kind in ('mismatches', 'missing', 'duplicates')}
    mismatch_sql, mismatch_params = list_query('mismatches', month)
    for label, condition in (('high_score', 'score >= ?'), ('low_score', 'score < ?')):
        write_csv(out / f'mismatches_{label}.csv', db.execute(
            f'SELECT * FROM ({mismatch_sql}) WHERE {condition}', [*mismatch_params, SCORE_THRESHOLD]))
    for label, status in (('no_digits', 'missing'), ('no_meter', 'no_meter'), ('blank', 'blank'),
                          ('no_saved_reading', 'no_saved_reading'), ('errors', 'error')):
        write_csv(out / f'{label}.csv', db.execute(
            f'SELECT month, {_PICKED}, error, read_at FROM readings WHERE status = ?' +
            (' AND month IS ?' if month else '') + ' ORDER BY picture_name',
            [status, _month(month)] if month else [status]))
    write_csv(out / 'reading_attempts.csv', db.execute(
        'SELECT * FROM reading_attempts' + (' WHERE month IS ?' if month else '') +
        ' ORDER BY read_at, picture_name', [_month(month)] if month else []))
    if full:  # one file per month: a whole year would not open in Excel
        months = [month] if month else [m for (m,) in db.execute(
            'SELECT DISTINCT month FROM readings ORDER BY month')]
        for month in months:
            month = month or 'unknown'
            write_csv(out / f'all_readings_{month}.csv', db.execute(*list_query('all', month)))
            # Every column of every photo, the model's whole answer and the image checks included.
            write_csv(out / f'details_{month}.csv', db.execute(
                'SELECT * FROM readings WHERE month IS ? ORDER BY picture_name',
                [None if month == 'unknown' else _month(month)]))

    if not quiet:
        print(f'{out}: summary.csv, ' + ', '.join(f'{kind}.csv ({rows})' for kind, rows in counts.items())
              + (', all_readings_<month>.csv' if full else '') + ', reading_attempts.csv')
        for line in (out / 'summary.csv').read_text(encoding='utf-8-sig').splitlines():
            print('  ' + line.replace(',', '\t'))


def command_report(args):
    db = connect(args.db)
    try:
        export_reports(db, args.out, args.month, args.full)
    finally:
        db.close()


class _LogTee:
    """Everything a run prints, each line starting with the time, written both to the
    terminal and to a log file, so the record survives a dropped connection."""

    def __init__(self, stream, log, lock):
        self.stream, self.log, self.lock, self.line_start = stream, log, lock, True

    def write(self, text):
        with self.lock:
            out = []
            for piece in text.splitlines(keepends=True):
                if self.line_start:
                    out.append(datetime.now().strftime('%Y-%m-%d %H:%M:%S '))
                out.append(piece)
                self.line_start = piece.endswith('\n')
            out = ''.join(out)
            self.stream.write(out)
            self.log.write(out)
        return len(text)

    def flush(self):
        self.stream.flush()
        self.log.flush()

    def __getattr__(self, name):  # encoding, isatty, ... of the terminal
        return getattr(self.stream, name)


def log_to_file(path):
    """Send stdout and stderr (tracebacks included) to the terminal and to path, timestamped."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    log, lock = path.open('a', encoding='utf-8', buffering=1), threading.Lock()
    sys.stdout = _LogTee(sys.stdout, log, lock)
    sys.stderr = _LogTee(sys.stderr, log, lock)
    print(f'Log file: {path.resolve()}', flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--db', default=str(AUDIT_DB))
    commands = parser.add_subparsers(dest='command', required=True)

    read = commands.add_parser('read', help='read photos and store the results')
    read.add_argument('folders', nargs='*')
    read.add_argument('--files-csv', help='read only paths in a CSV with picture_name and source_path; no folder scan')
    read.add_argument('--month', help='YYYY-MM for photos whose name carries no date')
    read.add_argument('--workers', type=int, help='photos read at once (default: from the model server)')
    read.add_argument('--model', choices=list(MODEL_BACKENDS), default='default',
                      help='model choice: default (FOTOMATES_OCR_URL) or m1, m2, ... from FOTOMATES_MODELS')
    read.add_argument('--limit', type=int, help='stop after this many photos (for a speed test)')
    read.add_argument('--skip-errors', action='store_true', help='do not retry photos that failed before')
    read.add_argument('--only-failed', action='store_true',
                      help='read again only the photos saved with a technical error ("Gabim teknik"), '
                           'nothing new; they are found by name, the folders are not listed')
    read.add_argument('--progress', type=int, default=100, help='print a progress line every N photos')
    read.add_argument('--report-every', type=int, default=1000,
                      help='refresh CSV reports after N newly processed photos (default: 1000; 0: only on exit)')
    read.add_argument('--verbose', action='store_true', help='print the file name, saved and AI readings, score and result for every photo')
    read.add_argument('--word-dir', default='reports', help='folder for the Word report of the run')
    read.add_argument('--no-word', action='store_true', help='do not write the Word report')
    read.add_argument('--log', help='log file for everything the run prints, each line timestamped '
                                    '(default: <word-dir>/logs/read_<date>_<time>.log)')
    read.add_argument('--no-log', action='store_true', help='print to the terminal only, write no log file')
    read.set_defaults(run=command_read)

    imported = commands.add_parser('import-log', help='store the results of a saved `read --verbose` output, '
                                                     'so those photos are not read again')
    imported.add_argument('logs', nargs='+', help='text file(s) with the terminal output of a read --verbose run')
    imported.add_argument('--month', help='YYYY-MM for photos whose name carries no date')
    imported.set_defaults(run=command_import_log)

    rematch = commands.add_parser('rematch', help='match the stored readings again, without reading photos')
    rematch.set_defaults(run=command_rematch)

    report = commands.add_parser('report', help='write the CSV reports')
    report.add_argument('--out', default='reports/audit')
    report.add_argument('--month', help='only this YYYY-MM')
    report.add_argument('--full', action='store_true', help='also export every row, one CSV per month')
    report.set_defaults(run=command_report)

    args = parser.parse_args()
    if args.command == 'read':
        if not args.folders and not args.files_csv:
            parser.error('read requires photo folders or --files-csv')
        if args.files_csv and (args.folders or args.only_failed):
            parser.error('--files-csv cannot be combined with folders or --only-failed')
        if not args.no_log:
            log_to_file(args.log or Path(args.word_dir) / 'logs' / f'read_{datetime.now():%Y%m%d_%H%M%S}.log')
            print('Command: ' + ' '.join(sys.argv), flush=True)
    args.run(args)


if __name__ == '__main__':
    main()
